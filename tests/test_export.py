"""Test §16.16: CSV and Excel exports work offline."""

import csv
import json

from taxengine import export as exportmod


def _make_assessment(test_dbs):
    from taxengine import decision, calc as calcmod
    k, w = test_dbs["knowledge"], test_dbs["workdata"]
    facts = {"tax_category": "withholding_tax", "jurisdiction": "federal",
             "taxpayer_type": "company", "transaction_type": "payment_for_goods",
             "transaction_date": "2025-02-01", "payment_category": "goods",
             "payment_amount": 200000, "payer_is_withholding_agent": True}
    steps = decision.analyze(
        k, {}, facts,
        calc_runner=lambda code, inputs, on_date=None:
            calcmod.run_calculation(k, code, inputs, on_date=on_date))
    aid = w.execute(
        "INSERT INTO assessments(title, facts) VALUES (?,?)",
        ("export test", json.dumps(facts))).lastrowid
    for s in steps:
        w.execute(
            "INSERT INTO assessment_steps(assessment_id, step_no, step_name,"
            " rule_code, finding, detail, confidence) VALUES (?,?,?,?,?,?,?)",
            (aid, s["step_no"], s["step_name"], "RULE-TEST-WHT",
             s["finding"], json.dumps(s["detail"]), s["confidence"]))
    w.execute("INSERT INTO calculation_runs(assessment_id, calc_code,"
              " calc_version, inputs, legislation_ref, effective_date, steps,"
              " rounding_method, result, validation_status)"
              " VALUES (?,?,?,?,?,?,?,?,?,?)",
              (aid, "CALC-TEST-WHT", 2, json.dumps({"payment_amount": 200000}),
               "Test Act s.1", "2024-07-01", json.dumps(["a", "b"]),
               "nearest_rupee", 14000.0, "complete"))
    w.execute("INSERT INTO evidence_tracker(assessment_id, evidence_code,"
              " requested, received) VALUES (?,?,1,1)", (aid, "EV-TEST-INV"))
    w.execute("INSERT INTO reconciliations(assessment_id, template_code,"
              " title, lines, total_difference, conclusion)"
              " VALUES (?,?,?,?,?,?)",
              (aid, "RECON-WHT-CREDITS", "t", json.dumps([]), 0.0, "ok"))
    w.commit()
    return aid


def test_rules_csv_export(test_dbs, tmp_path):
    k = test_dbs["knowledge"]
    p = exportmod.export_rules_csv(k, tmp_path / "rules.csv",
                                   status="verified")
    with open(p, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert rows and all(r["verification_status"] == "verified" for r in rows)
    assert any(r["code"] == "RULE-TEST-WHT" for r in rows)


def test_calculations_csv_export(test_dbs, tmp_path):
    w = test_dbs["workdata"]
    aid = _make_assessment(test_dbs)
    p = exportmod.export_calculations_csv(
        w, tmp_path / "calcs.csv", assessment_id=aid)
    with open(p, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 1
    assert rows[0]["calc_code"] == "CALC-TEST-WHT"
    assert float(rows[0]["result"]) == 14000.0


def test_assessment_xlsx_export(test_dbs, tmp_path):
    from openpyxl import load_workbook
    k, w = test_dbs["knowledge"], test_dbs["workdata"]
    aid = _make_assessment(test_dbs)
    p = exportmod.export_assessment_xlsx(k, w, aid, tmp_path / "a.xlsx")
    wb = load_workbook(p)
    assert {"Summary", "Rules", "Risks", "Evidence", "Calculations",
            "Reconciliations"} <= set(wb.sheetnames)
    summary = wb["Summary"]
    assert summary.max_row >= 15  # header + 14 steps
    logged = w.execute(
        "SELECT * FROM export_log WHERE assessment_id=?", (aid,)).fetchone()
    assert logged["file_format"] == "xlsx"
