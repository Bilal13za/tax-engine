"""Export module (Module H): CSV and Excel exports of rules, calculations,
risk assessments and reconciliations. Offline; uses stdlib csv + openpyxl.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path


def _rows_dicts(rows):
    return [dict(r) for r in rows]


def export_rules_csv(conn, path: Path, tax_category=None, jurisdiction=None,
                     status=None) -> Path:
    path = Path(path)
    sql = "SELECT code, version, title, tax_category, jurisdiction, " \
          "taxpayer_types, triggering_transaction, expected_treatment, " \
          "effective_from, effective_to, verification_status, reviewer, " \
          "review_date FROM tax_rules WHERE 1=1"
    params = []
    if tax_category:
        sql += " AND tax_category = ?"
        params.append(tax_category)
    if jurisdiction:
        sql += " AND jurisdiction = ?"
        params.append(jurisdiction)
    if status:
        sql += " AND verification_status = ?"
        params.append(status)
    sql += " ORDER BY code, version"
    rows = _rows_dicts(conn.execute(sql, params))
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()) if rows else [
            "code", "version", "title", "tax_category", "jurisdiction",
            "taxpayer_types", "triggering_transaction", "expected_treatment",
            "effective_from", "effective_to", "verification_status",
            "reviewer", "review_date"])
        w.writeheader()
        w.writerows(rows)
    return path


def export_calculations_csv(wconn, path: Path, assessment_id=None) -> Path:
    path = Path(path)
    sql = ("SELECT id, assessment_id, calc_code, calc_version, inputs, "
           "legislation_ref, effective_date, rounding_method, result, "
           "validation_status, created_at FROM calculation_runs")
    params = []
    if assessment_id:
        sql += " WHERE assessment_id = ?"
        params.append(assessment_id)
    rows = _rows_dicts(wconn.execute(sql, params))
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[
            "id", "assessment_id", "calc_code", "calc_version", "inputs",
            "legislation_ref", "effective_date", "rounding_method", "result",
            "validation_status", "created_at"])
        w.writeheader()
        w.writerows(rows)
    return path


def export_assessment_xlsx(kconn, wconn, assessment_id: int, path: Path) -> Path:
    """Full assessment workbook: summary of the 14 steps, rules, risks,
    evidence tracker, calculations, reconciliations."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E2F3")

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
        for r in rows:
            ws.append([str(v) if v is not None else "" for v in r])
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
        return ws

    a = wconn.execute("SELECT * FROM assessments WHERE id=?",
                      (assessment_id,)).fetchone()
    if a is None:
        raise KeyError(f"assessment {assessment_id} not found")
    a = dict(a)

    # Summary — the 14 steps
    steps = wconn.execute(
        "SELECT step_no, step_name, finding, confidence FROM assessment_steps "
        "WHERE assessment_id=? ORDER BY step_no", (assessment_id,)).fetchall()
    sheet("Summary",
          ["Assessment", a["title"], f"ID {assessment_id}", a["created_at"]],
          [[s["step_no"], s["step_name"], s["finding"], s["confidence"]]
           for s in steps])

    # Rules referenced (codes from workdata, details from knowledge base)
    rule_codes = [r["rule_code"] for r in wconn.execute(
        "SELECT DISTINCT rule_code FROM assessment_steps "
        "WHERE assessment_id=? AND rule_code IS NOT NULL", (assessment_id,))]
    rule_rows = []
    for rc in rule_codes:
        rr = kconn.execute(
            "SELECT code, version, title, tax_category, jurisdiction,"
            " verification_status FROM tax_rules WHERE code=? "
            "ORDER BY version DESC LIMIT 1", (rc,)).fetchone()
        if rr:
            rule_rows.append(rr)
    sheet("Rules", ["Code", "Version", "Title", "Category", "Jurisdiction", "Status"],
          [[r["code"], r["version"], r["title"], r["tax_category"],
            r["jurisdiction"], r["verification_status"]] for r in rule_rows])

    # Risks
    risks = wconn.execute(
        "SELECT step_no, finding, detail FROM assessment_steps "
        "WHERE assessment_id=? AND step_no=7", (assessment_id,)).fetchone()
    risk_rows = []
    if risks:
        for rk in json.loads(risks["detail"] or "{}").get("risks", []):
            risk_rows.append([rk.get("risk_code"), rk.get("title"),
                              (rk.get("description") or "")[:200],
                              rk.get("exposure_note") or ""])
    sheet("Risks", ["Risk code", "Title", "Description", "Exposure note"], risk_rows)

    # Evidence tracker
    ev = wconn.execute(
        "SELECT evidence_code, requested, received, verification_done, result, "
        "exceptions, outstanding_info, reviewer_comments, status "
        "FROM evidence_tracker WHERE assessment_id=?", (assessment_id,)).fetchall()
    sheet("Evidence",
          ["Evidence", "Requested", "Received", "Verified", "Result",
           "Exceptions", "Outstanding", "Reviewer comments", "Status"],
          [[e["evidence_code"], e["requested"], e["received"],
            e["verification_done"], e["result"], e["exceptions"],
            e["outstanding_info"], e["reviewer_comments"], e["status"]]
           for e in ev])

    # Calculations
    calcs = wconn.execute(
        "SELECT calc_code, calc_version, legislation_ref, effective_date, "
        "rounding_method, result, validation_status, steps "
        "FROM calculation_runs WHERE assessment_id=?", (assessment_id,)).fetchall()
    calc_rows = []
    for c in calcs:
        steps_txt = "\n".join(json.loads(c["steps"] or "[]"))
        calc_rows.append([c["calc_code"], c["calc_version"], c["legislation_ref"],
                          c["effective_date"], c["rounding_method"], c["result"],
                          c["validation_status"], steps_txt])
    ws = sheet("Calculations",
               ["Calc", "Version", "Provision", "Effective", "Rounding",
                "Result", "Validation", "Working"], calc_rows)
    for row in ws.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")

    # Reconciliations
    recs = wconn.execute(
        "SELECT template_code, title, lines, total_difference, conclusion "
        "FROM reconciliations WHERE assessment_id=?", (assessment_id,)).fetchall()
    rec_rows = []
    for rc in recs:
        for ln in json.loads(rc["lines"] or "[]"):
            rec_rows.append([rc["template_code"], rc["title"], ln.get("label"),
                             ln.get("amount_a"), ln.get("amount_b"),
                             ln.get("difference"), ln.get("status"),
                             ln.get("source_a"), ln.get("source_b")])
        rec_rows.append([rc["template_code"], rc["title"], "CONCLUSION: " +
                         (rc["conclusion"] or ""), "", "", rc["total_difference"],
                         "", "", ""])
    sheet("Reconciliations",
          ["Template", "Title", "Line", "Amount A", "Amount B", "Difference",
           "Status", "Source A", "Source B"], rec_rows)

    wb.save(path)
    wconn.execute(
        "INSERT INTO export_log(assessment_id, export_kind, file_format, file_name)"
        " VALUES (?,?,?,?)",
        (assessment_id, "full_assessment", "xlsx", path.name))
    wconn.commit()
    return path
